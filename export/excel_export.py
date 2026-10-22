"""openpyxl workbooks. All offline."""
import io
from datetime import date
import sqlite3

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


_HEADER_FILL = PatternFill("solid", fgColor="1F2937")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_THIN = Side(style="thin", color="D1D5DB")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _autosize(ws):
    for col in ws.columns:
        m = max((len(str(c.value)) for c in col if c.value is not None), default=0)
        ws.column_dimensions[col[0].column_letter].width = min(max(10, m + 3), 40)


def _style_header(ws):
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="left")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for cell in row:
            cell.border = _BORDER


def daily_xlsx(conn, for_date: str = None) -> bytes:
    d = for_date or date.today().isoformat()
    wb = Workbook()

    # Sheet 1: Sales
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Receipt", "Time", "Payment", "Cashier", "Items", "Total (GHS)"])
    rows = conn.execute("""
        SELECT s.receipt_no, s.created_at, s.payment, s.cashier_id,
               (SELECT COUNT(*) FROM sale_items si WHERE si.sale_id = s.id) AS items,
               s.total
        FROM sales s WHERE substr(s.created_at,1,10)=?
        ORDER BY s.id""", (d,)).fetchall()
    for r in rows:
        ws.append([r["receipt_no"], r["created_at"], r["payment"],
                   r["cashier_id"] or "", r["items"], float(r["total"])])
    _style_header(ws)
    _autosize(ws)

    # Sheet 2: Items
    ws = wb.create_sheet("Items")
    ws.append(["Product", "Qty", "Revenue (GHS)"])
    rows = conn.execute("""
        SELECT p.name, SUM(si.qty) AS qty, SUM(si.line_total) AS revenue
        FROM sale_items si
        JOIN sales s ON s.id = si.sale_id
        JOIN products p ON p.id = si.product_id
        WHERE substr(s.created_at,1,10)=?
        GROUP BY p.id ORDER BY revenue DESC""", (d,)).fetchall()
    for r in rows:
        ws.append([r["name"], int(r["qty"]), float(r["revenue"])])
    _style_header(ws)
    _autosize(ws)

    # Sheet 3: Summary
    ws = wb.create_sheet("Summary")
    ws.append(["Metric", "Value"])
    total = conn.execute(
        "SELECT COALESCE(SUM(total),0) FROM sales WHERE substr(created_at,1,10)=?",
        (d,)).fetchone()[0]
    n = conn.execute(
        "SELECT COUNT(*) FROM sales WHERE substr(created_at,1,10)=?",
        (d,)).fetchone()[0]
    ws.append(["Date", d])
    ws.append(["Total (GHS)", float(total)])
    ws.append(["Sales", n])
    for r in conn.execute(
        """SELECT payment, SUM(total) FROM sales
           WHERE substr(created_at,1,10)=? GROUP BY payment""", (d,)):
        ws.append([r[0].upper(), float(r[1])])
    _style_header(ws)
    _autosize(ws)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def stock_xlsx(conn) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Stock"
    ws.append(["Product", "Barcode", "Category", "Price (GHS)",
               "Cost (GHS)", "Stock", "Low at", "Expiry", "Active"])
    rows = conn.execute("""
        SELECT * FROM products ORDER BY name COLLATE NOCASE""").fetchall()
    for r in rows:
        ws.append([
            r["name"], r["barcode"] or "", r["category"] or "",
            float(r["price"]), float(r["cost"] or 0),
            int(r["stock"]), int(r["low_stock"] or 5),
            r["expiry_date"] or "", "Yes" if r["is_active"] else "No",
        ])
    _style_header(ws)
    _autosize(ws)

    # Highlight low-stock rows
    red = PatternFill("solid", fgColor="FEE2E2")
    for i, r in enumerate(rows, start=2):
        if r["is_active"] and r["stock"] <= (r["low_stock"] or 5):
            for cell in ws[i]:
                cell.fill = red

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()